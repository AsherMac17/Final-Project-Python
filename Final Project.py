from datetime import datetime
from openpyxl import Workbook
from openpyxl.chart import LineChart, BarChart, Reference

# The convertData function converts temperature from F to C.
def convertData(data): 
    convertedValue = (data -32) * 5 / 9
    return convertedValue

# The insertData function adds data to a CSV file
def insertData(filePath, data):
    try:
        with open(filePath, "a")as file:
            file.write(data + "\n")
        return True
    except Exception as error:
        print("Error while writing to the file:", error)
        return False

# The viewData function displays the contents of the CSV file and its path
def viewData(filePath):
    try:
        print("\nReading data from:", filePath)
        with open(filePath, "r") as file:
            print(file.read())
    except Exception as error:
        print("Error while reading the file:", error)

# The getInput function gets the user's input and saves each entry to the CSV file
def getInput():
    entries = int(input("How many entries are you inputting?: "))

    for i in range(entries):
        try:
            date = input("\nEnter a date: ")
            temperature = float(input("Enter the highest temp for the inputted date: "))
        
            convertedValue = convertData(temperature)

            data = date + "," + str(convertedValue)

            if insertData("ZooData.csv", data):
                print("The following data was saved at", datetime.now(), ":", data + ".")
        except Exception as error:
            print("Error while entering data:", error)

# The createChart function takes the CSV file path and chart type as arguments.
def createChart(filePath, chartType):
    try:
        source = input("Choose data source (1 = Celsius): ")

        if source != "1":
            print("Invalid data source.")
            return

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Temperature Data"

        sheet["A1"] = "Date"
        sheet["B1"] = "Celsius"

        with open(filePath, "r") as file:
            for rowNumber, line in enumerate(file, start=2):
                data = line.strip().split(",")

                if len(data) >= 2:
                    sheet.cell(row=rowNumber, column=1).value = data[0]
                    sheet.cell(row=rowNumber, column=2).value = float(data[1])

        if chartType == "line":
            chart = LineChart()
        else:
            chart = BarChart()

        chart.title = "ashmac9358 " + datetime.now().strftime("%m/%d/%Y")

        dates = Reference(
            sheet,
            min_col=1,
            min_row=2,
            max_row=sheet.max_row
        )

        values = Reference(
            sheet,
            min_col=2,
            min_row=1,
            max_row=sheet.max_row
        )

        chart.add_data(values, titles_from_data=True)
        chart.set_categories(dates)

        chart.x_axis.title = "Date"
        chart.x_axis.tickLblPos = "low"
        chart.x_axis.textRotation = 45

        chart.y_axis.title = "Temperature (Celsius)"
        chart.y_axis.tickLblPos = "nextTo"
        chart.y_axis.textRotation = 0

        chart.x_axis.delete = False
        chart.y_axis.delete = False

        chart.width = 18
        chart.height = 10

        chart.legend = None

        sheet.add_chart(chart, "D4")

        workbook.save("final.xlsx")

        print("\nReport created successfully!")
        print("Excel file: final.xlsx")
        print("Chart type:", chartType)
        print("Data source: Celsius")

    except Exception as error:
        print("Error creating report:", error)
        

# The generateReport function takes a CSV file path as an argument, and asks the user to choose a line/bar graph.
def generateReport(filePath):
    chartType = input("Choose graph type (line or bar): ").lower()

    if chartType == "line" or chartType == "bar":
        createChart(filePath, chartType)
    else:
        print("Invalid graph type.")

print("ashmac9358's Spreadsheet Automation Menu")
print("Choose a number from the following options")

menuOptions = ["1 Input Data", "2 View Current Data", "3 Generate Report"]

for option in menuOptions:
    print(option)

choice = input()

if choice == "1":
    getInput()
elif choice == "2":
    viewData("ZooData.csv")
elif choice == "3":
    generateReport("ZooData.csv")
else:
    print("Error: The chosen functionality is not implemented yet")

        
